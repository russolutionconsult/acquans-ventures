from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import sys

def main():
    document = Document()

    # Title
    title = document.add_heading('PROJECT CLOSEOUT & HANDOVER AGREEMENT', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Header Info
    document.add_paragraph("Date: May 29, 2026")
    document.add_paragraph("Project Name: Acquans Ventures Official Website Development")
    document.add_paragraph("Client: Acquans Ventures")
    document.add_paragraph("Developer/Agency: Russolution Consult")
    document.add_paragraph("Project URL: https://acquansventures.com")

    # Sections
    document.add_heading('1. Project Summary', level=1)
    document.add_paragraph("This document serves as the formal agreement and acknowledgment that the Acquans Ventures website project has been completed according to the agreed-upon scope of work, is fully deployed, and is now live for public access.")

    document.add_heading('2. Completed Deliverables', level=1)
    document.add_paragraph("The Developer has successfully delivered and implemented the following:")

    ul = document.add_paragraph(style='List Bullet')
    ul.add_run("Full Website Development: ").bold = True
    ul.add_run("A custom, fully responsive website built with modern web technologies (React, Tailwind CSS).")

    ul2 = document.add_paragraph(style='List Bullet')
    ul2.add_run("Core Pages: ").bold = True
    ul2.add_run("Overview (Home), About Us, Services, Projects, Gallery, Why Choose Us, Blog, and Contact pages.")

    ul3 = document.add_paragraph(style='List Bullet')
    ul3.add_run("Performance & Speed: ").bold = True
    ul3.add_run("Optimized codebase and assets for rapid loading times on both mobile and desktop devices.")

    ul4 = document.add_paragraph(style='List Bullet')
    ul4.add_run("Search Engine Optimization (SEO):").bold = True
    document.add_paragraph("Comprehensive local SEO integration targeting 'Ghana', Meta titles/descriptions, Open Graph tags, Schema.org Structured Data, and sitemap/robots.txt configured.", style='List Bullet 2')

    ul5 = document.add_paragraph(style='List Bullet')
    ul5.add_run("Analytics & Tracking:").bold = True
    document.add_paragraph("Google Analytics 4 (Measurement ID: G-J3TMY4NB5Y) fully integrated, and Google Search Console verified via DNS.", style='List Bullet 2')

    ul6 = document.add_paragraph(style='List Bullet')
    ul6.add_run("Live Deployment: ").bold = True
    ul6.add_run("The website has been successfully deployed to the production environment and is securely hosted with an active SSL certificate.")


    document.add_heading('3. Handover Materials', level=1)
    document.add_paragraph("Upon signing this closeout agreement, the Client assumes full ownership of the final website deliverables. The Developer has provided access to the Source code repository, Hosting platform, Google Analytics, Google Search Console, and Domain registrar.")

    document.add_heading('4. Post-Launch Support & Warranty', level=1)
    document.add_paragraph("The Developer agrees to a 30-day bug-fixing and warranty period starting from the date of this agreement. During this time, any critical errors or bugs found in the code will be fixed free of charge. New features requested after this date will be subject to a new agreement.")

    document.add_heading('5. Sign-off and Acceptance', level=1)
    document.add_paragraph("By signing below, Acquans Ventures acknowledges that the website has been reviewed, tested, and is accepted as complete and fully operational. The project is officially considered closed.")

    document.add_paragraph("\n")
    document.add_paragraph("For the Client (Acquans Ventures):").bold = True
    document.add_paragraph("Name: __________________________________")
    document.add_paragraph("Signature: _______________________________")
    document.add_paragraph("Date: ___________________________________")

    document.add_paragraph("\n")
    document.add_paragraph("For the Developer (Russolution Consult):").bold = True
    document.add_paragraph("Name: __________________________________")
    document.add_paragraph("Signature: _______________________________")
    document.add_paragraph("Date: ___________________________________")

    document.save('Project_Closeout_Agreement_Final.docx')

if __name__ == '__main__':
    main()
