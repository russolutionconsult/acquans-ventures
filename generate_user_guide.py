from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import sys

def main():
    document = Document()

    # Title
    title = document.add_heading('ACQUANS VENTURES - CLIENT PORTAL USER GUIDE', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Introduction
    document.add_paragraph("Welcome to your dedicated Client Portal! This platform allows you to monitor your projects in real-time and communicate directly with the Acquans Ventures management team.")

    # Section 1
    document.add_heading('1. Accessing Your Account', level=1)
    ul = document.add_paragraph(style='List Bullet')
    ul.add_run("Login: ").bold = True
    ul.add_run("Navigate to the login page. Enter the email address associated with your service request and the temporary password provided by our team.")
    
    ul2 = document.add_paragraph(style='List Bullet')
    ul2.add_run("Account Settings: ").bold = True
    ul2.add_run("Once logged in, you can view your registered email and update your contact name in the 'Account Settings' tab.")

    # Section 2
    document.add_heading('2. Navigating the Dashboard', level=1)
    document.add_paragraph("The dashboard is split into several main views, accessible from the sidebar menu:")
    
    ul3 = document.add_paragraph(style='List Bullet')
    ul3.add_run("Overview: ").bold = True
    ul3.add_run("Get a quick glance at your ongoing projects, handed-over (completed) projects, and any pending service requests.")
    
    ul4 = document.add_paragraph(style='List Bullet')
    ul4.add_run("Active Projects: ").bold = True
    ul4.add_run("View a detailed history of your projects. You can track execution progress (0-100%), view the start date, monitor site locations, and read the latest site status updates.")

    ul5 = document.add_paragraph(style='List Bullet')
    ul5.add_run("Communications: ").bold = True
    ul5.add_run("Access your real-time messaging threads. Unread message indicators will alert you to new replies directly in the sidebar.")

    # Section 3
    document.add_heading('3. Tracking Project Progress', level=1)
    document.add_paragraph("In the 'Active Projects' tab, each project card displays:")
    document.add_paragraph("A dynamic progress bar showing the verified site progress.", style='List Bullet')
    document.add_paragraph("The current project status (e.g., In Progress, Delivered).", style='List Bullet')
    document.add_paragraph("The latest site status update securely posted by the administrative team.", style='List Bullet')

    # Section 4
    document.add_heading('4. Communicating with the Team', level=1)
    document.add_paragraph("Click the 'Message Project Admin' button on any active project, or navigate to the 'Communications' tab to open your chats.", style='List Bullet')
    document.add_paragraph("This opens a dedicated chat interface for that specific project.", style='List Bullet')
    document.add_paragraph("You can send messages in real-time to the project managers.", style='List Bullet')
    document.add_paragraph("You do not need to refresh the page; new messages will appear automatically as they are sent!", style='List Bullet')

    # Section 5
    document.add_heading('5. Security & Logout', level=1)
    document.add_paragraph("When you are finished using the portal, ensure you click 'Sign Out' located at the bottom of the sidebar to secure your account session.", style='List Bullet')

    # Save
    document.save('Acquans_Ventures_Client_Portal_Guide.docx')
    print("User guide docx generated.")

if __name__ == '__main__':
    main()
